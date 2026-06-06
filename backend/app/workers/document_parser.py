"""
Document parser: extracts text content from PDF, DOCX, MD, TXT, Excel, and images.
Preserves structure metadata (page numbers, section titles) where possible.
Supports multi-encoding fallback for Chinese text files (UTF-8 → GB18030 → GBK).
"""

import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)

# Encoding fallback chain for Chinese text files
_TEXT_ENCODINGS = ["utf-8", "utf-8-sig", "gb18030", "gbk", "latin-1"]

_IMAGE_TYPES = frozenset({"png", "jpg", "jpeg", "webp", "gif", "bmp"})


def _read_text_file(file_path: str) -> str:
    """Read a text file with automatic encoding detection.

    Tries common encodings in order, falling back gracefully.
    Ensures Windows Chinese files (GBK/GB18030) are read correctly.
    """
    last_error = None
    for enc in _TEXT_ENCODINGS:
        try:
            with open(file_path, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError) as e:
            last_error = e
            continue
    raise ValueError(f"Failed to decode {file_path}: {last_error}")


def parse_file(file_path: str, file_type: str) -> list[dict]:
    """Parse a document file and return list of text segments.

    Args:
        file_path: Absolute path to the file.
        file_type: One of pdf, docx, md, txt, xlsx, xls, or image extensions.

    Returns:
        List of dicts: {"content": str, "page_number": int|None, "section_title": str|None}
    """
    ft = file_type.lower()
    parsers = {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "md": _parse_markdown,
        "txt": _parse_txt,
        "xlsx": _parse_xlsx,
        "xls": _parse_xls,
    }

    if ft in _IMAGE_TYPES:
        parser = _parse_image
    else:
        parser = parsers.get(ft)
    if not parser:
        raise ValueError(f"Unsupported file type: {file_type}")

    logger.info(f"Parsing {file_type} file: {file_path}")
    segments = parser(file_path)

    if not segments:
        logger.warning(f"No text content extracted from {file_path}")
        raise ValueError("No text content could be extracted from the document.")

    return segments


def _parse_pdf(file_path: str) -> list[dict]:
    """Parse PDF using PyMuPDF, preserving page numbers."""
    import fitz  # PyMuPDF

    doc = fitz.open(file_path)
    segments = []

    for page_num, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        if not text:
            continue

        # Try to extract a section title from the first meaningful line
        lines = [l.strip() for l in text.split("\n") if l.strip()]
        section_title = None
        if lines:
            # Heuristic: first short line (<=80 chars) might be a heading
            first_line = lines[0]
            if len(first_line) <= 80 and not first_line.endswith("."):
                section_title = first_line

        segments.append({
            "content": text,
            "page_number": page_num,
            "section_title": section_title,
        })

    doc.close()
    return segments


def _parse_docx(file_path: str) -> list[dict]:
    """Parse DOCX using python-docx, preserving heading structure."""
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)
    segments = []
    current_heading = None
    current_text: list[str] = []

    def flush():
        nonlocal current_text
        if current_text:
            content = "\n".join(current_text).strip()
            if content:
                segments.append({
                    "content": content,
                    "page_number": None,  # python-docx doesn't expose page numbers
                    "section_title": current_heading,
                })
            current_text = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        # Check if this is a heading
        if para.style.name.startswith("Heading"):
            flush()
            current_heading = text
            segments.append({
                "content": text,
                "page_number": None,
                "section_title": current_heading,
            })
            continue

        current_text.append(text)

    flush()

    # Also extract tables text
    for table in doc.tables:
        table_text: list[str] = []
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                table_text.append(row_text)
        if table_text:
            segments.append({
                "content": "\n".join(table_text),
                "page_number": None,
                "section_title": None,
            })

    return segments


def _parse_markdown(file_path: str) -> list[dict]:
    """Parse Markdown, extracting headings as section titles."""
    text = _read_text_file(file_path)

    segments = []
    current_heading = None
    current_lines: list[str] = []

    heading_pattern = re.compile(r"^(#{1,6})\s+(.+)$", re.MULTILINE)

    # Simple state-based parser
    lines = text.split("\n")
    for line in lines:
        match = heading_pattern.match(line)
        if match:
            # Flush previous segment
            if current_lines:
                content = "\n".join(current_lines).strip()
                if content:
                    segments.append({
                        "content": content,
                        "page_number": None,
                        "section_title": current_heading,
                    })
                current_lines = []

            current_heading = match.group(2)
            segments.append({
                "content": line,
                "page_number": None,
                "section_title": current_heading,
            })
        else:
            current_lines.append(line)

    # Flush remaining
    if current_lines:
        content = "\n".join(current_lines).strip()
        if content:
            segments.append({
                "content": content,
                "page_number": None,
                "section_title": current_heading,
            })

    return segments


def _parse_txt(file_path: str) -> list[dict]:
    """Parse plain text, splitting by paragraphs (double newlines)."""
    text = _read_text_file(file_path)

    segments = []
    paragraphs = re.split(r"\n\s*\n", text)

    for para in paragraphs:
        content = para.strip()
        if not content:
            continue

        lines = [l.strip() for l in content.split("\n") if l.strip()]
        section_title = None
        if lines and len(lines[0]) <= 80:
            section_title = lines[0]

        segments.append({
            "content": content,
            "page_number": None,
            "section_title": section_title,
        })

    return segments


def _parse_xlsx(file_path: str) -> list[dict]:
    """Parse XLSX workbook sheets into text segments."""
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    segments: list[dict] = []
    try:
        for sheet in wb.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                segments.append({
                    "content": "\n".join(rows),
                    "page_number": None,
                    "section_title": sheet.title,
                })
    finally:
        wb.close()
    return segments


def _parse_xls(file_path: str) -> list[dict]:
    """Parse legacy XLS workbook sheets into text segments."""
    import xlrd

    book = xlrd.open_workbook(file_path)
    segments: list[dict] = []
    for sheet in book.sheets():
        rows: list[str] = []
        for row_idx in range(sheet.nrows):
            cells: list[str] = []
            for col_idx in range(sheet.ncols):
                val = sheet.cell_value(row_idx, col_idx)
                if val is not None and str(val).strip():
                    cells.append(str(val).strip())
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            segments.append({
                "content": "\n".join(rows),
                "page_number": None,
                "section_title": sheet.name,
            })
    return segments


def _parse_image(file_path: str) -> list[dict]:
    """Extract image metadata and optional OCR text for RAG + preview."""
    from PIL import Image

    from app.services.ocr_service import extract_text_from_image

    path = Path(file_path)
    with Image.open(file_path) as img:
        width, height = img.size
        fmt = img.format or path.suffix.lstrip(".").upper()
        mode = img.mode

    lines = [
        f"图片文件：{path.name}",
        f"尺寸：{width} × {height}",
        f"格式：{fmt}",
        f"色彩模式：{mode}",
    ]

    ocr_text = extract_text_from_image(file_path)
    if ocr_text:
        lines.extend(["", "识别文字：", ocr_text])

    content = "\n".join(lines)
    return [{
        "content": content,
        "page_number": None,
        "section_title": path.name,
    }]
